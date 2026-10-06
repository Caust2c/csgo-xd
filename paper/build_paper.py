"""Build the measured research manuscript from the retained Word reference.

Dependencies: python-docx and Pillow. Run from the project root after run_all.R
and paper/export_paper_stats.R. --results supports a committed results snapshot.
The reference remains unchanged; unsupported headline claims are not inserted.
"""
from __future__ import annotations
import argparse
import csv
from copy import deepcopy
import hashlib
import json
import re
from pathlib import Path
import shutil
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageOps, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser()
parser.add_argument('--results', default='output')
parser.add_argument('--out', default='paper/CSGO_xD_Research_Paper.docx')
args = parser.parse_args()
RESULTS = ROOT / args.results
OUT = ROOT / args.out
REFERENCE = ROOT / 'paper/template/reference.docx'
EXPECTED_HASH = '505463a5a741fe08334e8c97d46bf3779a3d953d3cb0cdba98a57412b70b2be7'
if hashlib.sha256(REFERENCE.read_bytes()).hexdigest() != EXPECTED_HASH:
    raise ValueError('Reference changed; inspect the template before rebuilding.')

def read_csv(path):
    with path.open(encoding='utf-8-sig', newline='') as stream:
        return list(csv.DictReader(stream))

manifest = json.loads((RESULTS / 'run_manifest.json').read_text())
if manifest['smoke']:
    raise ValueError('The research paper requires the production results, not smoke outputs.')
metrics = read_csv(RESULTS / 'model_comparison.csv')
audit = read_csv(RESULTS / 'import_audit.csv')
support = read_csv(RESULTS / 'heatmap_support.csv')
descriptive = json.loads((ROOT / 'paper/inputs/descriptive_stats.json').read_text())
weapons = read_csv(ROOT / 'paper/inputs/weapon_summary.csv')
hits = read_csv(ROOT / 'paper/inputs/hitbox_summary.csv')
baseline = json.loads((RESULTS / 'baseline_metrics.json').read_text())
winner = manifest['selected_model']
chosen = next(r for r in metrics if r['model'] == winner)
raw_n = sum(int(r['raw_rows']) for r in audit)
clean_n = sum(int(r['clean_rows']) for r in audit)
sample_n = manifest['rows']
f = lambda r, key: float(r[key])
fraction = lambda n: 100 * int(n) / sample_n

OUT.parent.mkdir(parents=True, exist_ok=True)
shutil.copyfile(REFERENCE, OUT)
doc = Document(OUT)
original = Document(REFERENCE)
patterns = {name: deepcopy(original.paragraphs[index]._p) for name, index in
            [('body', 9), ('title', 0), ('abstract', 6), ('heading1', 8),
             ('heading2', 19), ('caption', 32), ('equation', 49)]}
body = doc._element.body
for child in list(body):
    if child.tag != qn('w:sectPr'):
        body.remove(child)

def paragraph(text='', pattern='body', size=None, bold=None, italic=None, keep=False):
    from docx.text.paragraph import Paragraph
    element = deepcopy(patterns[pattern])
    for child in list(element):
        if child.tag != qn('w:pPr'):
            element.remove(child)
    # Remove source separator rules and inherited highlighted placeholders.
    properties = element.find(qn('w:pPr'))
    if properties is not None:
        for border in list(properties.findall(qn('w:pBdr'))):
            properties.remove(border)
    body.insert(len(body) - 1, element)
    p = Paragraph(element, doc._body)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size or {'title': 17, 'abstract': 10, 'heading1': 12,
                               'heading2': 11, 'caption': 9, 'equation': 11}.get(pattern, 11))
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = bold if bold is not None else pattern in ('title', 'heading1', 'heading2')
    run.italic = italic if italic is not None else pattern in ('heading2', 'equation')
    p.paragraph_format.keep_with_next = keep or pattern in ('heading1', 'heading2')
    p.paragraph_format.widow_control = True
    if pattern == 'title':
        p.style = doc.styles['Title']
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p

def h1(text): return paragraph(text, 'heading1')
def h2(text): return paragraph(text, 'heading2')
def equation(text, number):
    from docx.enum.text import WD_TAB_ALIGNMENT
    p = paragraph('', 'equation')
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.tab_stops.clear_all()
    p.paragraph_format.tab_stops.add_tab_stop(Inches(3.1), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.2), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run('\t')
    for token in re.split(r'(_[a-z0-9]+|\^[0-9]+)', text):
        run = p.add_run(token[1:] if token.startswith(('_', '^')) else token)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
        run.italic = True
        if token.startswith('_'): run.font.subscript = True
        if token.startswith('^'): run.font.superscript = True
    p.add_run(f'\t({number})').font.size = Pt(11)
    return p

table_count = 0
def table(title, headers, rows, widths):
    global table_count
    table_count += 1
    paragraph(f'Table {table_count}. {title}', 'caption', keep=True)
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    grid = t._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        element = OxmlElement('w:gridCol')
        element.set(qn('w:w'), str(int(width * 1440)))
        grid.append(element)
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{edge}')
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            border.set(qn('w:' + key), value)
        borders.append(border)
    props.append(borders)
    margin = OxmlElement('w:tblCellMar')
    for edge in ['top', 'left', 'bottom', 'right']:
        node = OxmlElement('w:' + edge)
        node.set(qn('w:w'), '65')
        node.set(qn('w:type'), 'dxa')
        margin.append(node)
    props.append(margin)
    all_rows = [headers] + list(rows)
    for index, values in enumerate(all_rows):
        row = t.rows[0] if index == 0 else t.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if index == 0: trpr.append(OxmlElement('w:tblHeader'))
        for j, value in enumerate(values):
            cell = row.cells[j]
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(str(value))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(9)
            run.bold = index == 0
            run.font.color.rgb = RGBColor(0, 0, 0)
            if index == 0:
                fill = OxmlElement('w:shd')
                fill.set(qn('w:fill'), 'D9E2F3')
                cell._tc.get_or_add_tcPr().append(fill)
    paragraph('', size=4).paragraph_format.space_after = Pt(3)
    return t

figure_count = 0
def figure(path, caption, width=6.1):
    global figure_count
    figure_count += 1
    p = paragraph('', keep=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    image = p.add_run().add_picture(str(path), width=Inches(width))
    image._inline.docPr.set('descr', caption)
    cp = paragraph(f'Fig. {figure_count}. {caption}', 'caption')
    cp.paragraph_format.keep_with_next = False
    return p

assets = ROOT / 'paper/figures'
assets.mkdir(parents=True, exist_ok=True)

def montage(paths, target, labels, cols=2, tile=(1000, 970)):
    rows = (len(paths) + cols - 1) // cols
    band = 45 if labels else 0
    sheet = Image.new('RGB', (cols * tile[0], rows * (tile[1] + band)), 'white')
    draw = ImageDraw.Draw(sheet)
    try:
        font = ImageFont.truetype('C:/Windows/Fonts/timesbd.ttf', 28)
    except OSError:
        font = ImageFont.load_default()
    for i, path in enumerate(paths):
        im = Image.open(path).convert('RGB')
        im = ImageOps.contain(im, tile)
        x, y = (i % cols) * tile[0], (i // cols) * (tile[1] + band)
        sheet.paste(im, (x + (tile[0] - im.width) // 2, y + band))
        if labels: draw.text((x + 20, y + 5), labels[i], fill='black', font=font)
    sheet.save(target)
    return target

paragraph('Expected Damage Modelling and Spatial Diagnostics in Counter Strike Global Offensive', 'title')
for index in range(1, 5):
    body.insert(len(body) - 1, deepcopy(original.paragraphs[index]._p))
paragraph('', size=4)
abstract = (
    'Abstract - This study evaluates conditional expected damage (xD) in Counter-Strike: Global Offensive using a modular R workflow that combines disk-backed data retrieval, regression benchmarking and spatial diagnostics. '
    f'Both ESEA damage parts of the public Kaggle archive supply {raw_n:,} raw events; cleaning retains {clean_n:,}. A map-balanced sample of {sample_n:,} events from 240 matches limits model memory, with a shared 20,000-row fitting subset and disjoint validation and test matches. '
    'Twelve regression methods are compared using training-only encoding and feature selection. Validation selects XGBoost, which attains test MAE '
    f'{f(chosen,"MAE"):.3f}, RMSE {f(chosen,"RMSE"):.3f} and R squared {f(chosen,"R2"):.3f} on 8,000 events from 32 held-out matches. '
    f'The match-bootstrap 95% MAE interval is {f(chosen,"MAE_lower"):.3f} to {f(chosen,"MAE_upper"):.3f}. Its MAE improvement over a weapon-hitbox mean baseline is {f(chosen,"MAE_improvement"):.3f}, with a confidence interval crossing zero and Holm-adjusted p-value {f(chosen,"p_holm"):.4f}. '
    'Model-specific radar heatmaps mask cells with insufficient event and match support. Separate static and interactive graphics document the sampled event distribution. Results support a reproducible conditional damage baseline and descriptive spatial diagnostics; they do not establish superior prediction over the simple baseline, tactical causation or calibrated expected kills.'
)
paragraph(abstract, 'abstract')
paragraph('Keywords - Counter-Strike Global Offensive, expected damage, regression benchmarking, SQLite, match-level validation, radar heatmaps, reproducible R analysis', 'abstract', italic=True)

h1('I INTRODUCTION')
paragraph('Combat telemetry records where players meet, which weapons they use and how much damage each observed hit delivers. Aggregate kills and total damage summarize outcomes, but they do not separate observed damage from the expectation implied by weapon, hit location and context. This distinction motivates a conditional expected-damage baseline: an event delivering 30 damage may be ordinary for one combination of weapon and hitbox and unusual for another.')
paragraph('Expected-goals models in football estimate the probability that a shot becomes a goal from its context [1]. We borrow the comparison between realized and expected outcomes while changing the target to a continuous quantity. The analogy has an important boundary. Our records already condition on a damage-producing event and include the realized hitbox; they cannot estimate a pre-shot kill probability or value every opportunity available to a player.')
paragraph('The project addresses three research questions. RQ1 examines the distributions of damage, weapon class and timing in a documented analysis sample. RQ2 compares 12 regression methods against a simple training-only weapon-hitbox baseline on matches excluded from fitting. RQ3 asks how predicted damage and residuals can be displayed spatially without treating sparsely observed locations as reliable tactical findings.')
paragraph('The practical contribution is a reproducible research workflow rather than a claim of a new learning algorithm. It imports both damage parts into SQLite, retrieves a bounded sample, freezes training-derived preprocessing, compares learner families, quantifies match-level uncertainty and renders model-specific held-out heatmaps. Resource controls target a 16 GB laptop setting by bounding chunk sizes, sample sizes, model complexity and worker threads.')

h1('II RELATED WORK')
h2('A Action valuation and player ratings')
paragraph('Xenopoulos, Doraiswamy and Silva developed a context-aware framework that values CS:GO actions through changes in a team\'s chance of winning and includes a data model and graph-based distance representation [2]. This addresses action consequence at the round level. Our xD target instead describes the amount of damage in a recorded event, without claiming that equal damage has equal round-winning value.')
paragraph('Xu and Moka study player plus/minus ratings using regularized, logistic and Bayesian models to relate participation to team point differences [3]. Their player-level outcome objective differs from this event-level conditional damage objective. An xD residual therefore cannot be interpreted directly as a player contribution rating.')
h2('B Regression methods and comparison design')
paragraph('Coordinate-descent regularization paths provide efficient ridge, lasso and elastic-net estimators [4]. Tree ensembles accommodate nonlinear relationships: ranger implements regression forests efficiently [5], while XGBoost provides a scalable tree-boosting system [6]. These methods motivate part of our benchmark, alongside a regression tree, MARS, a generalized additive model and a shallow neural network. Statistical comparison remains necessary because model complexity alone does not demonstrate an improvement over a context baseline.')
h2('C Position of the present study')
paragraph('The benchmark complements prior player and action valuation by examining a simpler retrospective outcome with explicit resource limits. Its central question is whether additional spatial and temporal predictors improve materially over weapon and hitbox information. The recorded-hit selection mechanism, held-out match protocol and baseline comparison are treated as part of the model definition, not as incidental implementation details.')

h1('III DATASET')
h2('A Archive and analysis population')
paragraph('The data are the CS:GO Competitive Matchmaking Data archive hosted on Kaggle [7]. Production ingestion uses esea_master_dmg_demos.part1.csv and part2.csv, with their corresponding metadata files to associate each demo filename with a map. The archive also contains matchmaking, grenade and kill tables. Those tables have different schemas and are not silently pooled into the damage model.')
table('Production ingestion counts', ['Damage part', 'Raw events', 'Retained events'],
      [[r['source'].replace('esea_master_dmg_demos.', ''), f"{int(r['raw_rows']):,}", f"{int(r['clean_rows']):,}"] for r in audit] +
      [['Total', f'{raw_n:,}', f'{clean_n:,}']], [2.8, 1.7, 1.7])
paragraph(f'Cleaning retains {clean_n/raw_n*100:.2f}% of raw rows. The retained population consists of positive, finite hostile-player damage with usable coordinates and known context. It does not include misses, unrecorded opportunities or all environmental damage. Modeling uses {sample_n:,} sampled events rather than all retained rows; full-database SQL summaries and sample-based charts are labelled separately.')
h2('B Fields and schema validation')
table('Required input fields and derived variables', ['Fields', 'Role', 'Interpretation'], [
    ['file, round, tick', 'Grouping', 'Match identity and within-match event ordering; not model predictors.'],
    ['hp_dmg, arm_dmg', 'Outcome', 'Health and armor damage components; excluded from predictors.'],
    ['wp, wp_type, hitbox', 'Context', 'Observed weapon, class and impact location.'],
    ['att_side, vic_side', 'Filtering/context', 'Hostile-player filtering; victim side becomes redundant.'],
    ['seconds, bomb state/site', 'Timing/context', 'Parser time and bomb context.'],
    ['att/vic X and Y', 'Spatial', 'Planar world coordinates for distance and map plotting.'],
    ['map, distance, log distance, time bin', 'Derived', 'Metadata association and engineered predictors.']], [2.25, 1.1, 2.85])
paragraph('CSV schemas are checked before streaming. Metadata must provide a unique map for each filename within its part. Numeric coercion is explicit; missing damage is dropped rather than imputed as zero. Player identifiers, ranks and team names are not stored as model predictors or in the exported spatial caches.')

h1('IV METHODOLOGY')
h2('A Modular workflow and database connectivity')
paragraph('Configuration, ingestion, preprocessing, fitting, statistics, visualization and reporting are implemented as separate R scripts and functions. SQLite provides actual persistent database storage and retrieval through DBI and RSQLite [8]. Import callbacks parse only the required columns in 50,000-row chunks and append cleaned rows inside a per-file transaction. A filename index supports parameterized retrieval; SQL also computes full-import map and weapon-class aggregates.')
table('Core scripts and responsibilities', ['Component', 'Responsibility'], [
    ['config.R and installer', 'Portable paths, bounded settings and dependency checks.'],
    ['database.R', 'Chunked cleaning, metadata association, SQLite import and SQL summaries.'],
    ['pipeline.R', 'Sample retrieval, match splits, shared fitting subset and model comparison.'],
    ['R/preprocess.R', 'Training-only encoding, redundancy filtering and scaling.'],
    ['R/models.R and statistics.R', 'Uniform learner adapters, tuning and match-level uncertainty.'],
    ['heatmaps.R', 'Supported held-out spatial means and radar overlays.'],
    ['visualizations.R and report.R', 'Separate EDA/Plotly figures, tables, measured report and gallery.'],
    ['verify_results.R', 'Leakage, metrics, model reload and spatial support checks.']], [2.0, 4.2])
paragraph('The SQLite cache is approximately 64 MiB and temporary tables use disk. Completed imports can be reused; an interrupted part rolls back. File-size and modification-time checks detect stale inputs but are not cryptographic integrity guarantees. If source data or cleaning rules change, the database must be rebuilt. The full production database is about 2.17 GB on disk; that size is not a RAM measurement.')
h2('B Cleaning and feature engineering')
paragraph('Both players must belong to opposing CounterTerrorist/Terrorist sides. Each damage component must be finite and within 0 to 100, total damage must be positive, and time, round, tick and planar positions must be finite. Coordinates at the paired zero sentinel are removed. Missing/unknown weapon classes and invalid hitbox codes are excluded; absent bomb-site strings become none. These rules define the study population and may exclude unusual legitimate records.')
equation('y_i = H_i + A_i', 1)
equation('d_i = √[(x_a − x_v)^2 + (y_a − y_v)^2]', 2)
equation('xD_i = estimated E[y_i | observed hit, z_i]', 3)
paragraph('Here H and A are health and armor damage, a and v denote attacker and victim, and z is the available event context.')
paragraph('Engineering adds planar distance, log1p(distance), a numeric bomb indicator and time bins at 20 and 60 seconds. These are parser-time bins, because the meaning of seconds is not independently verified as a tactical phase. Map identity is included so that identical world coordinates on different maps can be distinguished. Victim side is omitted from predictors because hostile-player filtering determines it from attacker side.')
h2('C Sampling and match separation')
paragraph('A seeded round robin balances sampled matches across maps; small maps exhaust their available supply. Within each selected match, SQL uses a deterministic tick/round pseudo-random order and retains at most 250 rows. This is a bounded sampling device, not a proven uniform event sampler. All sampled events of a match share one partition. Splits are approximately 70/15/15 within maps when at least three matches are available; tiny map groups receive seeded probability assignments.')
table('Observed analysis partitions', ['Partition', 'Sampled matches', 'Rows used'], [
    ['Training partition', descriptive['train_matches'], f"{sample_n-manifest['validation_rows']-manifest['test_rows']:,} available; 20,000 fit"],
    ['Validation', descriptive['validation_matches'], f"{manifest['validation_rows']:,}"],
    ['Testing', descriptive['test_matches'], f"{manifest['test_rows']:,}"],
    ['Total sample', descriptive['sample_matches'], f'{sample_n:,}']], [2.1, 1.5, 2.6])
paragraph('A random capped subset of training-partition rows is shared by all learners. Validation determines hyperparameters and the selected method; test matches are used for reporting only. Assertions require disjoint filenames across partitions, nonempty partitions and at least five test matches. Seed 42 is fixed for sampling, fitting and uncertainty calculations.')
h2('D Training only feature selection')
paragraph('Categorical predictors are one-hot encoded using at most 25 training levels per field, with one reference level omitted. Rare or unseen levels receive the reference representation. Zero-variance columns are removed, and numeric predictors with absolute training correlation above .98 are removed sequentially. Remaining columns are standardized using training means and standard deviations. No damage outcome or identifier enters the feature matrix.')
paragraph(f'The production run retains {descriptive["selected_feature_count"]} encoded predictors. The frozen category levels, selected columns and scaling parameters are saved for inference. Validation permutation importance shuffles one encoded feature at a time for the selected learned model and measures the MAE increase. It is a diagnostic, not causal attribution or an additional test-driven selection step.')
h2('E Learners and validation tuning')
table('Twelve implemented regression methods', ['Method', 'R implementation', 'Validation choice or bound'], [
    ['Ordinary least squares', 'stats::lm.fit', 'Fixed linear fit; aliased coefficients set to zero.'],
    ['Ridge', 'glmnet, alpha 0', 'Lambda chosen on fitted regularization path.'],
    ['Lasso', 'glmnet, alpha 1', 'Lambda chosen on fitted regularization path.'],
    ['Elastic net', 'glmnet, alpha .5', 'Lambda chosen on fitted regularization path.'],
    ['CART', 'rpart', 'Complexity .005 or .02; depth at most 8.'],
    ['Random forest', 'ranger', 'Node size 10 or 30; 150 trees, depth at most 12.'],
    ['Extra trees', 'ranger extratrees', 'Node size 10 or 30; same forest size bound.'],
    ['Gradient boosting', 'gbm', 'Depth 2 or 4; 200 trees.'],
    ['XGBoost', 'xgboost', 'Depth 3 or 5; at most 250 rounds; patience 20.'],
    ['MARS', 'earth', 'Degree 1 or 2; at most 40 candidate basis terms.'],
    ['Generalized additive model', 'mgcv::bam', 'Small smooth bases; exact prediction after discretized fit.'],
    ['Neural network', 'nnet', 'One hidden layer with 5 or 10 units; 150 iterations.']], [1.45, 1.45, 3.3])
paragraph('The benchmark includes related regularized linear and forest methods; these are 12 methods, not 12 unrelated model families. The neural network is shallow. A weapon-by-hitbox mean computed on the same training subset is an additional baseline; unseen combinations use the global training mean. The baseline remains eligible for validation selection.')
paragraph('All tuning minimizes validation MAE. XGBoost uses a histogram tree method, learning rate .05, row and column subsampling .8 and at most two threads. Small candidate grids limit tuning cost. GAM fitting uses discretization for efficiency, but prediction is exact so that the same row receives the same score alone or in a larger batch. Nonfinite predictions or failed serialization are reported explicitly rather than counted as successful models.')
h2('F Statistical validation')
equation('MAE = (1/N) Σ_i |y_i − xD_i|', 4)
equation('RMSE = √[(1/N) Σ_i (y_i − xD_i)^2]', 5)
equation('R^2 = 1 − Σ_i (y_i − xD_i)^2 / Σ_i (y_i − ȳ)^2', 6)
paragraph('Signed bias is the mean of actual minus predicted damage. Test metrics are event-weighted on the same 8,000 rows for every method. Confidence intervals use 500 test-match bootstrap replicates, resampling matches with replacement and combining per-match sufficient statistics. This preserves sampled within-match dependence and avoids constructing large resampled event tables. Intervals are conditional on the fitted model and sample policy; they exclude uncertainty from refitting.')
paragraph('For each model, the paired difference is baseline absolute error minus model absolute error. Positive mean improvement favors the model. A match-level sign-flip randomization test uses 500 replicates and Holm adjustment across successful learners. The test assumes exchangeability or symmetry under the null and has finite Monte Carlo resolution. It is an exploratory comparison, not a causal claim or a replacement for external validation.')
h2('G Spatial aggregation and radar alignment')
equation('Δ_i = y_i − xD_i; mean residual in cell b = (1/n_b) Σ Δ_i', 7)
paragraph('Only held-out test predictions enter the spatial caches. Every represented map uses a fixed 40 by 40 grid in attacker coordinates. Mean layers require at least five events and two distinct matches; unsupported cells are omitted. Count layers retain all occupied cells. Grid CSVs preserve event count, match count, actual mean, predicted mean, residual mean and centers. Count maps show sampled recorded hits, not occupancy or exposure.')
paragraph('Radar bounds come from the supplied StartX/EndX/StartY/EndY calibration with independent axis scales. Increasing game Y points upward. Maps without a calibration entry and image are explicitly shown as world-coordinate schematics. Out-of-bounds events are counted and omitted from the display. Actual/predicted damage shares a 0-200 display range; residuals share -100 to 100. Display saturation does not alter numeric evaluation or CSV values.')

h1('V RESULTS')
h2('A Sample distributions and full database context')
rifle = next(r for r in weapons if r['wp_type'] == 'Rifle')
pistol = next(r for r in weapons if r['wp_type'] == 'Pistol')
paragraph(f'In the map-balanced sample, rifles account for {int(rifle["events"]):,} events ({fraction(rifle["events"]):.2f}%) and pistols for {int(pistol["events"]):,} ({fraction(pistol["events"]):.2f}%). The overall median damage is {descriptive["damage_median"]:.0f}, mean {descriptive["damage_mean"]:.2f}, and observed range {descriptive["damage_min"]:.0f} to {descriptive["damage_max"]:.0f}. These are exact sample summaries, not population frequency estimates.')
table('Damage by weapon class in the bounded sample', ['Class', 'Events', 'Median', 'Interquartile range'],
      [[r['wp_type'], f"{int(r['events']):,}", f"{float(r['median']):g}", f"{float(r['q25']):g} to {float(r['q75']):g}"] for r in weapons],
      [1.5, 1.45, 1.0, 2.25])
figure(RESULTS / 'eda/weapon_class_counts.png', 'Weapon-class event counts in the bounded map-balanced analysis sample.', width=5.6)
paragraph('Weapon-class medians vary substantially: the sniper median is 100 damage points, rifle 29, pistol 25 and grenade 6. This context dependence explains why a weapon-hitbox baseline is informative. The figures describe event amounts and composition; they do not reveal how often players shoot, miss or occupy a location.')
figure(RESULTS / 'eda/damage_histogram.png', 'Health-plus-armor damage distribution for the analysis sample.', width=5.6)
paragraph(f'The most frequent sampled hitbox is {hits[0]["hitbox"]} with {int(hits[0]["events"]):,} events. Parser-time boxplots and range-versus-damage scatterplots are exported separately. Their associations should not be interpreted as phase or distance causing a damage change; weapon, hitbox and missing context can confound both. SQL map-count figures summarize all cleaned imported rows, while sample plots remain explicitly labelled.')
h2('B Predictive comparison on unseen matches')
metric_rows = [[r['model'], f"{f(r,'validation_MAE'):.3f}", f"{f(r,'MAE'):.3f}",
                f"{f(r,'RMSE'):.3f}", f"{f(r,'R2'):.3f}"] for r in metrics]
metric_rows.append(['Weapon-hitbox baseline', f"{baseline['validation_MAE']:.3f}",
                    f"{baseline['metrics']['MAE']:.3f}", f"{baseline['metrics']['RMSE']:.3f}", f"{baseline['metrics']['R2']:.3f}"])
table('Validation selection and test performance', ['Method', 'Val MAE', 'Test MAE', 'Test RMSE', 'Test R squared'],
      metric_rows, [2.0, 1.05, 1.05, 1.05, 1.05])
paragraph(f'All 12 methods completed. Validation MAE selects {winner} at {f(chosen,"validation_MAE"):.3f}; its test MAE is {f(chosen,"MAE"):.3f}, RMSE {f(chosen,"RMSE"):.3f}, R squared {f(chosen,"R2"):.3f} and signed bias {f(chosen,"Bias"):.3f}. The positive bias indicates a small average underprediction. The model is selected from validation results rather than whichever test row happens to have the smallest metric.')
figure(RESULTS / 'eda/model_comparison.png', 'Test MAE and match-bootstrap 95 percent intervals. The dashed line is the training-only weapon-hitbox baseline.', width=6.1)
paragraph(f'The selected model\'s MAE improvement over the baseline is {f(chosen,"MAE_improvement"):.3f} damage points, about {100*f(chosen,"MAE_improvement")/baseline["metrics"]["MAE"]:.2f}% of baseline MAE. Its 95% improvement interval is {f(chosen,"improvement_lower"):.3f} to {f(chosen,"improvement_upper"):.3f}, and Holm-adjusted p-value is {f(chosen,"p_holm"):.4f}. The interval includes zero. Accordingly, this run does not establish an improvement over the baseline. R squared near .70 describes conditional damage fit, not tactical intelligence.')
figure(RESULTS / 'eda/regression_calibration.png', 'Regression calibration of the validation-selected model by predicted-damage decile on test matches.', width=5.6)
paragraph('Calibration is assessed as mean realized versus mean predicted damage within test prediction deciles. This is regression calibration, not probability calibration. The residual plot and validation permutation-importance chart are included in the companion results gallery; they help locate remaining errors and dependence on features but do not identify causal mechanisms.')
h2('C Spatial diagnostics and support')
selected_support = [r for r in support if r['model'] == winner]
table('Held-out map support for the selected model', ['Map', 'Test events', 'Supported cells', 'Radar calibrated'],
      [[r['map'], r['events'], r['supported_cells'], 'Yes' if r['calibrated'] == 'TRUE' else 'No'] for r in selected_support],
      [1.65, 1.2, 1.55, 1.8])
paragraph('The run produces 224 heatmaps across eight test-represented maps: predicted and residual layers for 12 methods plus the baseline, and shared actual/count layers. Each map/model also has a CSV grid. A spatially empty patch can mean no recorded events or inadequate support; it is not evidence of zero danger. Radar calibration is supplied for only some maps, and planar views cannot separate floors.')
map_name = 'de_dust2'
panels = assets / 'panels'
panel = montage([panels / f'{winner}_expected.png', panels / f'{winner}_actual.png',
                 panels / f'{winner}_residual.png', panels / f'{winner}_events.png'],
                assets / 'dust2_spatial_layers.png', None)
figure(panel, 'Dust2 held-out spatial layers for XGBoost: xD, actual damage, residual and sampled event count. Unsupported mean cells are omitted.', width=6.2)
paragraph('Warm residual cells indicate greater realized damage than the fitted context baseline, and cool cells indicate lower damage. Near-zero residual cells indicate agreement with the baseline. Their interpretation remains conditional on recorded hits and the selected sample. Missing armor, remaining health and penetration can explain residuals; no cell is asserted to be a causal tactical advantage or a measure of player skill.')

h1('VI DISCUSSION')
h2('A What the comparison establishes')
paragraph('The study demonstrates that a complete multi-method workflow can analyze the archive without forming a full damage matrix in memory. Nonlinear and linear learners reach broadly similar conditional damage error scales. The strongest practical conclusion from this run is the adequacy of a simple context baseline, rather than a statistically supported advantage of the validation-selected learner. More complex models may still be useful for spatial diagnostics, but that use must be justified separately from aggregate fit.')
h2('B Limitations and threats to validity')
paragraph('Recorded-hit selection excludes misses and unobserved exposure. Hitbox is a realized impact feature, so the model describes retrospective damage rather than pre-shot effectiveness. The target combines health and armor damage and is not equivalent to kills, round value or remaining opponent threat. A large conditional R squared can mainly reflect familiar weapon and hitbox mechanics.')
paragraph('The map-balanced match sample and per-match row cap change event frequencies relative to the full archive. Tick/round ordering is deterministic but not proven uniform. Rare categories fall back to a shared reference representation, and some maps have limited support. One split and small validation grids do not establish robustness across seeds, time periods, map versions, matchmaking populations or CS2.')
paragraph('Victim armor, helmet, pre-hit health, wall penetration and height are missing. Distance is planar. Cell means have a support threshold but no cell-level confidence interval or multiple-testing correction. Radar overlays require landmark validation before tactical use; unsupported areas and multi-floor geometry constrain interpretation. Permutation importance can be distorted by correlated or dependent encoded variables.')
paragraph('Match-bootstrap intervals and paired randomization describe uncertainty conditional on the fitted model and sample. They do not include training-set refitting uncertainty or repeated-experiment selection effects. Native memory was not profiled: the 16 GB design relies on bounded work sizes, not an experimentally established peak-RAM guarantee. The completed run used Windows R 4.6.1; portable Linux instructions are provided, but native Linux execution is not claimed.')

h1('VII REPRODUCIBILITY AND COURSE COVERAGE')
h2('A Execution and validation evidence')
paragraph('The public companion repository is https://github.com/Caust2c/csgo-xd. Open csgo_xD.Rproj with the extracted archive beside it, install dependencies with source("00_install_packages.R"), and execute source("run_all.R"). CSGO_DATA_DIR overrides the archive location. A separate smoke mode exercises a prefix subset; the paper reports production-import results with bounded model sampling, not smoke values. The README contains Linux system-library guidance and lower-cost settings.')
paragraph('The repository includes the modular source, measured result snapshot, paper, figure gallery and saved settings. Large raw CSVs, SQLite databases and platform-specific libraries are regenerated locally rather than committed. The manifest records seed, sample policy, row counts, threads, bootstrap count and package versions. Package versions are recorded but not locked; an environment lockfile on the target platform would improve long-term replication.')
paragraph('Verification checks filename separation, exclusion of damage outcomes from features, finite predictions, independent metric recalculation, interval order, multiplicity adjustment, saved-model reload equivalence and spatial support invariants. All 12 saved models pass in a fresh R process. Syntax parsing covers the R workflow. Figure review corrected unsupported-bin opacity, and model reload review corrected GAM batch-dependent prediction discretization.')
h2('B Syllabus and evaluation criteria')
table('Implemented course and evaluation evidence', ['Area', 'Concrete implementation'], [
    ['Functions, variables and comments', 'Documented modular functions; shared config; conditionals and source loading.'],
    ['Vectors and lists', 'Filtering and lapply summaries exported independently from heatmaps.'],
    ['Data frames and CSV wrangling', 'Schema checks, chunk cleaning, metadata association and saved tables.'],
    ['dplyr and tidyr', 'Grouped summaries, joins, pivot_longer and pivot_wider demonstrations.'],
    ['Databases and web APIs', 'SQLite ingestion/retrieval; separate optional public REST/JSON example.'],
    ['Static and interactive graphics', '12 EDA PNGs and 3 Plotly HTML charts; radar mean/count layers.'],
    ['Feature engineering and selection', 'Distance/timing/bomb features; train-only encoding and redundancy filtering.'],
    ['10 to 15 ML or DL algorithms', '12 completed regression methods and an additional baseline.'],
    ['Statistical validation and reflection', 'Match splits, bootstrap intervals, paired tests, limits and xK roadmap.']], [2.1, 4.1])

h1('VIII CONCLUSION AND FUTURE WORK')
paragraph(f'This study provides a modular R benchmark and spatial diagnostic workflow for recorded CS:GO damage. SQLite ingestion retains {clean_n:,} events from both archive parts, while capped sampling supports shared fitting and evaluation on a laptop-oriented configuration. All 12 methods complete. Validation-selected XGBoost reaches test MAE {f(chosen,"MAE"):.3f} and R squared {f(chosen,"R2"):.3f}, but its advantage over the weapon-hitbox baseline is uncertain. Held-out map layers and separate EDA communicate where events are observed and how conditional errors vary, with explicit support restrictions.')
paragraph('Future work should define and label genuine kill opportunities before estimating xK. Reparsed demos would need player linkage, lethal/nonlethal outcomes, misses, exposure, health before impact, armor and height. Such a model should report log loss, Brier score, discrimination and probability calibration on match and temporal holdouts. Further extensions include grouped cross-validation, refit uncertainty, calibrated floor-specific geometry, per-cell match-bootstrap intervals and full streaming inference into SQL aggregates. None of those extensions is claimed as completed here.')

h1('REFERENCES')
references = [
    'G. Anzer and P. Bauer, "A Goal Scoring Probability Model for Shots Based on Synchronized Positional and Event Data in Football (Soccer)," Frontiers in Sports and Active Living, vol. 3, article 624475, 2021. doi: 10.3389/fspor.2021.624475.',
    'P. Xenopoulos, H. Doraiswamy and C. Silva, "Valuing Player Actions in Counter-Strike: Global Offensive," 2020. arXiv:2011.01324. https://arxiv.org/abs/2011.01324.',
    'H. Xu and S. Moka, "Rating Players of Counter-Strike: Global Offensive Based on Plus/Minus value," 2024. arXiv:2409.05052. https://arxiv.org/abs/2409.05052.',
    'J. Friedman, T. Hastie and R. Tibshirani, "Regularization Paths for Generalized Linear Models via Coordinate Descent," Journal of Statistical Software, vol. 33, no. 1, pp. 1-22, 2010. doi: 10.18637/jss.v033.i01.',
    'M. N. Wright and A. Ziegler, "ranger: A Fast Implementation of Random Forests for High Dimensional Data in C++ and R," Journal of Statistical Software, vol. 77, no. 1, pp. 1-17, 2017. doi: 10.18637/jss.v077.i01.',
    'T. Chen and C. Guestrin, "XGBoost: A Scalable Tree Boosting System," 2016. doi: 10.1145/2939672.2939785. https://arxiv.org/abs/1603.02754.',
    'skihikingkevin, "CS:GO Competitive Matchmaking Data," Kaggle. https://www.kaggle.com/datasets/skihikingkevin/csgo-matchmaking-damage. Accessed 6 October 2026.',
    'RSQLite contributors, "Connect to an SQLite database," RSQLite documentation. https://rsqlite.r-dbi.org/reference/SQLite.html. Accessed 6 October 2026.'
]
for i, ref in enumerate(references, 1):
    p = paragraph(f'[{i}] {ref}', size=9.5)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(.25)
    p.paragraph_format.first_line_indent = Inches(-.25)
    p.paragraph_format.space_after = Pt(5)

h1('APPENDIX MODEL HEATMAP COMPARISON')
paragraph('The following panels show Dust2 expected-damage means for all 12 methods on identical held-out events and support-filtered cells. They are descriptive comparisons under the same 0-200 color range. Full-resolution images, all eight represented maps and corresponding residual layers are available in the results gallery. Similar color patterns can reflect the shared strong weapon/hitbox dependence; they do not establish that a particular map position causes a higher damage outcome.')
for batch, start in enumerate(range(0, len(metrics), 4), 1):
    subset = metrics[start:start + 4]
    names = [r['model'] for r in subset]
    panel = montage([panels / f'{name}_expected.png' for name in names],
                    assets / f'dust2_algorithms_{batch}.png', None)
    figure(panel, 'Dust2 expected-damage comparison for ' + ', '.join(names) + '. Same held-out events, cell support and display scale.', width=6.2)

doc.core_properties.title = 'Expected Damage Modelling and Spatial Diagnostics in Counter Strike Global Offensive'
doc.core_properties.subject = 'Measured CS GO damage regression and spatial analysis'
doc.core_properties.author = 'Hrishikesh Damodar Nayak and Hardik Omesh Lalla'
doc.core_properties.last_modified_by = 'Hrishikesh Damodar Nayak and Hardik Omesh Lalla'
doc.core_properties.keywords = 'CSGO, expected damage, R, SQLite, regression, heatmaps'
doc.save(OUT)

# Preserve untouched opaque/package parts byte-for-byte from the retained source.
editable = {'word/document.xml', 'word/_rels/document.xml.rels', 'word/styles.xml',
            'docProps/core.xml', '[Content_Types].xml'}
temp = OUT.with_suffix('.tmp.docx')
with ZipFile(REFERENCE) as source, ZipFile(OUT) as built, ZipFile(temp, 'w', ZIP_DEFLATED) as merged:
    source_names = set(source.namelist())
    for info in built.infolist():
        data = source.read(info.filename) if info.filename in source_names and info.filename not in editable else built.read(info.filename)
        merged.writestr(info, data)
    for info in source.infolist():
        if info.filename not in built.namelist(): merged.writestr(info, source.read(info.filename))
temp.replace(OUT)

# Audit noneditable source package parts and geometry after authoring.
with ZipFile(REFERENCE) as source, ZipFile(OUT) as final:
    changed = [n for n in source.namelist() if n not in editable and source.read(n) != final.read(n)]
    if changed: raise ValueError('Preserve-only package parts changed: ' + repr(changed))
new_doc = Document(OUT)
section = new_doc.sections[0]
base_section = original.sections[0]
for field in ['page_width', 'page_height', 'left_margin', 'right_margin', 'top_margin', 'bottom_margin']:
    if getattr(section, field) != getattr(base_section, field): raise ValueError('Geometry changed: ' + field)
text = '\n'.join(p.text for p in new_doc.paragraphs) + '\n'.join(c.text for t in new_doc.tables for row in t.rows for c in row.cells)
for placeholder in ['[MAE]', '[RMSE]', '[R2]', '[Insert', '[add volume', 'optional:']:
    if placeholder in text: raise ValueError('Unresolved placeholder: ' + placeholder)
if hashlib.sha256(REFERENCE.read_bytes()).hexdigest() != EXPECTED_HASH: raise ValueError('Source template was modified.')
print(f'Created {OUT.name}: {len(new_doc.paragraphs)} paragraphs, {len(new_doc.tables)} tables, {len(new_doc.inline_shapes)} figures.')
print('Template geometry, source hash, preserve-only parts and placeholder checks passed.')
